"""Regression coverage for safety-critical commercial UI state transitions."""

from __future__ import annotations

import html
import json
from io import BytesIO
from pathlib import Path
from string import Formatter

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from shieldlab.physics import barriers, solver, sources
from shieldlab.regulatory import limits
from shieldlab.report import report as calculator_report
from shieldlab.room import diagram, report_regulatory, report_room
from shieldlab.room.decision_support import summarize_results
from shieldlab.room.engines import AnalyticalEngine, EngineResult, SurrogateEngine
from shieldlab.room.model import Opening, RoomDesign
from ui import i18n, room_results
from ui.room_results import RoomAssessment, row_status_for


ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "app.py"
ROOM_APP_PATH = ROOT / "pages" / "1_Room_Designer.py"


def _run_app(path: Path) -> AppTest:
    app = AppTest.from_file(str(path), default_timeout=60).run()
    assert not app.exception
    return app


def _element_with_label(elements, label: str):
    return next(element for element in elements if element.label == label)


def _calculator_with_three_distinct_layers() -> tuple[AppTest, tuple]:
    app = _run_app(APP_PATH)
    _element_with_label(app.button, "Add Layer").click().run()
    _element_with_label(app.button, "Add Layer").click().run()
    layer_values = (
        ("concrete", 111.0),
        ("lead", 2.0),
        ("steel", 33.0),
    )
    for index, (material, thickness) in enumerate(layer_values, start=1):
        _element_with_label(app.selectbox, f"Layer {index} Material").select(material)
        _element_with_label(
            app.number_input, f"Layer {index} Thickness (mm)"
        ).set_value(thickness)
    app.run()
    return app, layer_values


def _room_context_html(app: AppTest) -> str:
    return next(
        element.proto.body
        for element in app.get("html")
        if "sl-context-strip" in element.proto.body
        and "Engine State" in element.proto.body
    )


def _context_html(app: AppTest) -> str:
    return next(
        element.proto.body
        for element in app.get("html")
        if "sl-context-strip" in element.proto.body
    )


def _status_html(app: AppTest) -> str:
    return next(
        element.proto.body
        for element in app.get("html")
        if "sl-status-card" in element.proto.body
    )


def _rendered_markup(app: AppTest) -> str:
    markdown = [element.value for element in app.markdown]
    html = [element.proto.body for element in app.get("html")]
    return "\n".join(markdown + html)


def _format_fields(template: str) -> set[str]:
    return {
        field_name
        for _, field_name, _, _ in Formatter().parse(template)
        if field_name is not None
    }


def _engine_result(
    wall_id: str,
    *,
    dose: float,
    goal: float,
    transmission: float,
    ci_high: float,
) -> EngineResult:
    return EngineResult(
        barrier_id=wall_id,
        label=f"Wall {wall_id}",
        engine="surrogate",
        B_required=transmission,
        B_achieved=transmission,
        dose_mSv_wk=dose,
        goal_over_T=goal,
        passes=True,
        margin=goal / dose,
        material="concrete",
        ci_low=transmission * 0.8,
        ci_high=ci_high,
    )


def _ci_crossing_assessment() -> tuple[RoomAssessment, EngineResult, EngineResult]:
    critical = _engine_result(
        "N",
        dose=0.5,
        goal=1.0,
        transmission=0.1,
        ci_high=0.12,
    )
    noncritical_ci_crossing = _engine_result(
        "E",
        dose=0.25,
        goal=1.0,
        transmission=0.1,
        ci_high=0.5,
    )
    results = [critical, noncritical_ci_crossing]
    design = RoomDesign.default()
    assessment = RoomAssessment(
        design,
        "check",
        results,
        results,
        results,
        SurrogateEngine(design),
    )
    return assessment, critical, noncritical_ci_crossing


def _pdf_text(pdf_bytes: bytes) -> str:
    return "\n".join(
        page.extract_text() or ""
        for page in PdfReader(BytesIO(pdf_bytes)).pages
    )


def test_calculator_removing_middle_layer_preserves_remaining_values():
    app, layer_values = _calculator_with_three_distinct_layers()
    remove_buttons = [button for button in app.button if button.label == "Remove"]
    remove_buttons[1].click().run()

    remaining_layers = app.session_state["layers"]
    assert [
        (layer["material"], layer["thickness"]) for layer in remaining_layers
    ] == [layer_values[0], layer_values[2]]
    assert _element_with_label(app.selectbox, "Layer 2 Material").value == "steel"
    assert _element_with_label(
        app.number_input, "Layer 2 Thickness (mm)"
    ).value == 33.0
    assert not app.exception


def test_calculator_reset_discards_previous_layer_widget_state():
    app, _ = _calculator_with_three_distinct_layers()
    _element_with_label(app.button, "Reset Layers").click().run()

    reset_layers = app.session_state["layers"]
    assert len(reset_layers) == 1
    assert reset_layers[0]["material"] == "concrete"
    assert reset_layers[0]["thickness"] == 150.0
    assert _element_with_label(app.selectbox, "Layer 1 Material").value == "concrete"
    assert _element_with_label(
        app.number_input, "Layer 1 Thickness (mm)"
    ).value == 150.0
    assert not app.exception


def test_empty_room_json_is_rejected_without_replacing_current_design():
    app = _run_app(ROOM_APP_PATH)
    original_design = json.loads(app.session_state["design"].to_json())

    app.file_uploader(key="load_json").set_value(
        ("empty-room.json", b"{}", "application/json")
    )
    app.run()
    _element_with_label(app.button, "Load Selected Design").click().run()

    assert not app.exception
    assert any(
        "not a valid ShieldLab room design" in error.value
        for error in app.error
    )
    assert json.loads(app.session_state["design"].to_json()) == original_design


def test_loaded_room_preserves_custom_occupancy_and_zero_lead_equivalence():
    design = RoomDesign.default()
    north_wall = design.wall("N")
    north_wall.adjacent.occupancy_T = 0.25
    north_wall.openings.append(
        Opening(
            kind="door",
            center_along_wall_m=2.0,
            width_m=1.0,
            lead_equiv_mm=0.0,
        )
    )
    expected_design = json.loads(design.to_json())

    app = _run_app(ROOM_APP_PATH)
    app.file_uploader(key="load_json").set_value(
        ("custom-room.json", design.to_json().encode("utf-8"), "application/json")
    )
    app.run()
    _element_with_label(app.button, "Load Selected Design").click().run()

    assert not app.exception
    loaded_design = app.session_state["design"]
    assert loaded_design.wall("N").adjacent.occupancy_T == 0.25
    assert loaded_design.wall("N").openings[0].lead_equiv_mm == 0.0
    assert json.loads(loaded_design.to_json()) == expected_design
    occupancy_control = _element_with_label(app.selectbox, "Occupancy Factor (T)")
    assert occupancy_control.value == "Custom Project Value"
    custom_occupancy = _element_with_label(
        app.number_input, "Custom Occupancy Factor (T)"
    )
    assert custom_occupancy.value == 0.25
    lead_control = _element_with_label(app.number_input, "Lead Equivalence (mm)")
    assert lead_control.value == 0.0


def test_shrinking_room_with_opening_keeps_editor_renderable_and_geometry_valid():
    app = _run_app(ROOM_APP_PATH)
    app.button(key="add_opening_N").click().run()
    assert len(app.session_state["design"].wall("N").openings) == 1

    width_control = _element_with_label(
        app.number_input, "Width · West to East (m)"
    )
    width_control.set_value(1.0).run()

    assert not app.exception
    design = app.session_state["design"]
    opening = design.wall("N").openings[0]
    assert opening.width_m <= design.room.width_m
    assert opening.width_m / 2 <= opening.center_along_wall_m
    assert opening.center_along_wall_m + opening.width_m / 2 <= design.room.width_m
    assert not [error for error in design.validate() if "opening" in error.lower()]


def test_room_context_uses_source_values_changed_in_the_same_rerun():
    app = _run_app(ROOM_APP_PATH)
    _element_with_label(app.selectbox, "Radionuclide").select("I-131")
    _element_with_label(app.number_input, "Patients per Week").set_value(57.0)
    app.run()

    context_html = _room_context_html(app)
    assert 'title="I-131"' in context_html
    assert 'title="57 patients/week"' in context_html
    assert not app.exception


def test_noncritical_ci_crossing_forces_room_review_and_amber_plan_status():
    assessment, critical, noncritical_ci_crossing = _ci_crossing_assessment()
    results = assessment.decision_results

    summary = summarize_results(results)
    assert summary["critical"] is critical
    assert summary["status"] == "PASS"
    assert assessment.status == "REVIEW"
    assert any(
        "95% upper bound" in reason and "Wall E" in reason
        for reason in assessment.review_reasons
    )
    assert row_status_for(noncritical_ci_crossing) == ("review", "!", "Review")

    status_by_label = {
        result.label: row_status_for(result)[0]
        for result in assessment.decision_results
    }
    rendered_plan = diagram.render(
        assessment.design,
        results,
        status_by_label=status_by_label,
    )
    image = Image.open(BytesIO(rendered_plan)).convert("RGB")
    rendered_colors = {
        color
        for _, color in image.getcolors(maxcolors=image.width * image.height)
    }
    assert (194, 123, 0) in rendered_colors


def test_ncrp_room_exports_keep_goal_in_mgy_and_disclose_photon_approximation():
    assessment, _, _ = _ci_crossing_assessment()
    report = room_results._build_room_report(assessment)

    assert report["units"] == {"dose": "mSv/week", "goal": "mGy/week"}
    assert "1 mGy ≈ 1 mSv" in report["unit_note"]

    detailed_html = report_room.to_html(report).decode("utf-8")
    regulatory_html = report_regulatory.build_submission_html(
        report,
        {"facility": "UI regression facility"},
    ).decode("utf-8")
    for exported_html in (detailed_html, regulatory_html):
        assert "mGy/week" in exported_html
        assert report["unit_note"] in exported_html
        assert "Limit P/T (mSv/wk)" not in exported_html

    assert "Limit P/T (mGy/" in regulatory_html
    workbook = load_workbook(BytesIO(report_room.to_xlsx(report)), read_only=True)
    xlsx_headers = [cell.value for cell in workbook["Barriers"][1]]
    workbook.close()
    assert "limit P/T (mGy/week)" in xlsx_headers
    assert "limit P/T (mSv/week)" not in xlsx_headers

    summary_text = " ".join(_pdf_text(report_room.to_summary_pdf(report)).split())
    goal_section = summary_text.split("Regulatory design goal", 1)[1].split(
        "Safety margin", 1
    )[0]
    assert "mGy/week" in goal_section
    assert "mSv/week" not in goal_section
    assert "Unit basis" in summary_text
    assert "1 mGy" in summary_text and "1 mSv" in summary_text


def test_room_html_export_escapes_notes_loaded_from_project_json():
    # Regression 2026-08-10: project JSON notes previously rendered as executable markup.
    malicious_notes = '<img src=x onerror="alert(\'اختراق\')">'
    design = RoomDesign.default()
    design.notes = malicious_notes
    loaded_design = RoomDesign.from_json(design.to_json())
    analytical_results = AnalyticalEngine(loaded_design).evaluate_all("design")
    report = report_room.build_report(
        loaded_design,
        analytical_results,
        "design",
        b"",
    )

    exported_html, mime_type, extension = report_room.export(report, "HTML")
    rendered_html = exported_html.decode("utf-8")

    assert mime_type == "text/html"
    assert extension == "html"
    assert malicious_notes not in rendered_html
    assert html.escape(malicious_notes) in rendered_html


def test_ci_crossing_room_exports_remain_review_required():
    assessment, _, noncritical_ci_crossing = _ci_crossing_assessment()
    report = room_results._build_room_report(assessment)

    assert report["summary"]["status"] == "REVIEW"
    assert report["summary"]["review_reasons"] == assessment.review_reasons
    crossing_row = next(
        row for row in report["rows"]
        if row["barrier"] == noncritical_ci_crossing.label
    )
    assert crossing_row["verdict"] == "REVIEW"
    assert any(
        "95% upper bound" in reason
        for reason in report["summary"]["review_reasons"]
    )

    regulatory_html = report_regulatory.build_submission_html(
        report,
        {"facility": "UI regression facility"},
    ).decode("utf-8")
    assert "Overall assessment: REVIEW" in regulatory_html
    assert "REVIEW REQUIRED" in regulatory_html
    assert "not a declaration of compliance" in regulatory_html
    assert "shielding specified meets the stated design goals" not in regulatory_html

    summary_text = " ".join(_pdf_text(report_room.to_summary_pdf(report)).split())
    assert "REVIEW" in summary_text
    assert "REVIEW REQUIRED" in summary_text


def test_ncrp_calculator_pdf_labels_regulatory_goal_with_goal_unit():
    source = sources.radionuclide_point_source(
        "F-18",
        activity_mCi=10.0,
        d_m=2.0,
        hours_per_week=1.0,
    )
    barrier = barriers.Barrier().add("concrete", 150.0)
    provisional_goal = limits.design_goal("NCRP", "controlled", occupancy_T=1.0)
    provisional = solver.evaluate(source, barrier, provisional_goal)
    goal = limits.design_goal(
        "NCRP",
        "controlled",
        occupancy_T=1.0,
        override_P_weekly=provisional.transmitted_total * 1.1,
    )
    evaluation = solver.evaluate(source, barrier, goal)
    assert source.unit == "mSv/week"
    assert goal.unit == "mGy/week"
    assert 1.0 < evaluation.verdict.margin_ratio < 1.2

    pdf_bytes = calculator_report.build_pdf_summary(
        source=source,
        barrier=barrier,
        goal=goal,
        evaluation=evaluation,
        inputs={},
        prepared_by="",
        facility="",
    )
    pdf_text = " ".join(_pdf_text(pdf_bytes).split())
    assert "REVIEW REQUIRED" in pdf_text
    goal_section = pdf_text.split("Regulatory design goal / T", 1)[1].split(
        "Safety margin", 1
    )[0]
    assert f"{goal.P_weekly / goal.occupancy_T:.4g} {goal.unit}" in goal_section
    assert source.unit not in goal_section

def test_locale_catalogs_have_matching_nonempty_messages_and_placeholders():
    english_messages = i18n.MESSAGES["en"]
    arabic_messages = i18n.MESSAGES["ar"]

    assert english_messages.keys() == arabic_messages.keys()
    for message_key in english_messages:
        assert english_messages[message_key].strip()
        assert arabic_messages[message_key].strip()
        assert _format_fields(english_messages[message_key]) == _format_fields(
            arabic_messages[message_key]
        )


@pytest.mark.parametrize(
    ("language", "direction"),
    [("en", "ltr"), ("ar", "rtl")],
)
def test_language_toggle_renders_matching_direction_and_calculator_copy(
    language: str,
    direction: str,
):
    app = _run_app(APP_PATH)
    app.sidebar.segmented_control(key="_sl_language_control").select(language).run()

    context_html = _context_html(app)
    tab_labels = [tab.label for tab in app.tabs]
    assert app.sidebar.segmented_control(key="_sl_language_control").value == language
    assert f'lang="{language}"' in context_html
    assert f'dir="{direction}"' in context_html
    assert i18n.MESSAGES[language]["calculator_title"] in _rendered_markup(app)
    assert i18n.MESSAGES[language]["tab_assessment"] in tab_labels
    assert not app.exception


def test_arabic_language_survives_room_navigation_and_json_load():
    app = _run_app(APP_PATH)
    app.sidebar.segmented_control(key="_sl_language_control").select("ar").run()

    app.switch_page("pages/1_Room_Designer.py").run()
    assert app.session_state[i18n.LANGUAGE_KEY] == "ar"
    assert i18n.MESSAGES["ar"]["room_title"] in _rendered_markup(app)
    assert 'lang="ar"' in _context_html(app)
    assert 'dir="rtl"' in _context_html(app)

    room_design = RoomDesign.default()
    app.file_uploader(key="load_json").set_value(
        (
            "bilingual-room.json",
            room_design.to_json().encode("utf-8"),
            "application/json",
        )
    )
    app.run()
    app.button(key="load_selected_design").click().run()

    assert app.session_state[i18n.LANGUAGE_KEY] == "ar"
    assert i18n.MESSAGES["ar"]["room_title"] in _rendered_markup(app)
    assert not app.exception

    app.switch_page("app.py").run()
    assert app.session_state[i18n.LANGUAGE_KEY] == "ar"
    assert i18n.MESSAGES["ar"]["calculator_title"] in _rendered_markup(app)
    assert not app.exception


def test_language_switch_preserves_calculator_inputs_layers_and_decision():
    app = _run_app(APP_PATH)
    app.number_input(key="calc_diag_patients").set_value(321).run()
    app.button(key="calc_add_layer").click().run()

    first_layer, second_layer = app.session_state["layers"]
    app.number_input(key=f"premium_thk_{first_layer['_id']}").set_value(275.0).run()
    app.selectbox(key=f"premium_mat_{second_layer['_id']}").select("lead").run()
    app.number_input(key=f"premium_thk_{second_layer['_id']}").set_value(2.5).run()

    expected_layers = tuple(
        (layer["_id"], layer["material"], layer["thickness"])
        for layer in app.session_state["layers"]
    )
    expected_metrics = tuple(metric.value for metric in app.metric)
    english_status = _status_html(app)
    decision_key = next(
        message_key
        for message_key in ("pass", "review_required", "fail")
        if f">{i18n.MESSAGES['en'][message_key]} ·" in english_status
    )

    app.sidebar.segmented_control(key="_sl_language_control").select("ar").run()

    actual_layers = tuple(
        (layer["_id"], layer["material"], layer["thickness"])
        for layer in app.session_state["layers"]
    )
    assert app.session_state["calc_diag_patients"] == 321
    assert actual_layers == expected_layers
    assert tuple(metric.value for metric in app.metric) == expected_metrics
    assert (
        f">{i18n.MESSAGES['ar'][decision_key]} ·"
        in _status_html(app)
    )
    assert not app.exception








